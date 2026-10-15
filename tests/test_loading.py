"""Loading: both Excel sheets are stacked without losing rows, with the expected columns and types."""
import openpyxl
import pandas as pd
import pytest

from src.load_data import RAW_PATH, load_raw

from conftest import needs_raw

SHEETS = {"Year 2009-2010": 525_461, "Year 2010-2011": 541_910}   # reports/data_quality.md
TOTAL_ROWS = 1_067_371                                            # README, cleaning_log.md
EXCEL_COLUMNS = ["Invoice", "StockCode", "Description", "Quantity", "InvoiceDate",
                 "Price", "Customer ID", "Country"]
EXPECTED_DTYPES = {
    "Invoice": "string", "StockCode": "string", "Description": "string", "Quantity": "int64",
    "Price": "float64", "Customer ID": "Int64", "Country": "string", "source_sheet": "category",
}


def check_dtypes(df: pd.DataFrame) -> None:
    for col, dtype in EXPECTED_DTYPES.items():
        assert str(df[col].dtype) == dtype, f"{col}: {df[col].dtype}, expected {dtype}"
    assert pd.api.types.is_datetime64_dtype(df["InvoiceDate"])


# ---------- load_raw on a tiny workbook (always runs) ----------

@pytest.fixture
def tiny_workbook(tmp_path):
    path = tmp_path / "tiny.xlsx"
    sheet1 = pd.DataFrame({
        "Invoice": ["489434", "C489449"], "StockCode": ["85048", "22087"],
        "Description": ["GLASS BALL", 1234],            # a number in a text column, as in the real file
        "Quantity": [12, -12], "InvoiceDate": pd.to_datetime(["2009-12-01 07:45", "2009-12-01 10:33"]),
        "Price": [6.95, 2.95], "Customer ID": [13085.0, None], "Country": ["United Kingdom", "Germany"],
    })
    sheet2 = sheet1.iloc[[0]].assign(Invoice="539993", InvoiceDate=pd.Timestamp("2011-01-04 10:00"))
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        sheet1.to_excel(xl, sheet_name="Year 2009-2010", index=False)
        sheet2.to_excel(xl, sheet_name="Year 2010-2011", index=False)
    return path


def test_load_raw_stacks_all_sheets(tiny_workbook):
    df = load_raw(tiny_workbook)
    assert len(df) == 3
    assert df["source_sheet"].value_counts().to_dict() == {"Year 2009-2010": 2, "Year 2010-2011": 1}
    assert list(df.columns) == EXCEL_COLUMNS + ["source_sheet"]


def test_load_raw_sets_types(tiny_workbook):
    df = load_raw(tiny_workbook)
    check_dtypes(df)
    assert df.loc[1, "Description"] == "1234"                 # numeric cell read as text
    assert df["Customer ID"].isna().sum() == 1                 # missing ID kept as <NA>, not dropped
    assert df.loc[0, "Customer ID"] == 13085


# ---------- the real Excel file and the loaded raw data ----------

@needs_raw
def test_excel_file_sheets_headers_and_row_counts():
    """Read the workbook structure only (read-only mode), without loading every cell."""
    wb = openpyxl.load_workbook(RAW_PATH, read_only=True)
    try:
        assert wb.sheetnames == list(SHEETS)
        for name, rows in SHEETS.items():
            ws = wb[name]
            header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            assert header == EXCEL_COLUMNS
            assert ws.max_row - 1 == rows, f"{name}: {ws.max_row - 1:,} data rows"
    finally:
        wb.close()


def test_raw_row_count(raw):
    assert len(raw) == TOTAL_ROWS
    assert raw["source_sheet"].value_counts().to_dict() == SHEETS


def test_raw_columns(raw):
    assert list(raw.columns) == EXCEL_COLUMNS + ["source_sheet"]


def test_raw_dtypes(raw):
    check_dtypes(raw)


def test_raw_date_range(raw):
    assert raw["InvoiceDate"].min() == pd.Timestamp("2009-12-01 07:45:00")
    assert raw["InvoiceDate"].max() == pd.Timestamp("2011-12-09 12:50:00")
